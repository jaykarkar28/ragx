
"""Document loaders for the RAGX ingestion pipeline."""

from dataclasses import dataclass, field
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


class DocumentLoadError(Exception):
    """Raised when a document cannot be loaded or parsed."""


@dataclass
class LoadedDocument:
    """Standard representation of an extracted document."""

    content: str
    metadata: dict = field(default_factory=dict)


class DocumentLoader:
    """Load text content from supported document formats."""

    def load_file(self, file_path: str | Path) -> list[LoadedDocument]:
        """
        Load a single document.

        PDFs produce one LoadedDocument per non-empty page.
        Other formats produce one LoadedDocument per file.
        """
        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(f"File not found: {path}")

        if not path.is_file():
            raise ValueError(f"Expected a file, got: {path}")

        extension = path.suffix.lower()

        if extension not in SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Unsupported file format: {extension}. "
                f"Supported formats: {sorted(SUPPORTED_EXTENSIONS)}"
            )

        try:
            if extension == ".pdf":
                documents = self._load_pdf(path)
            elif extension == ".docx":
                documents = self._load_docx(path)
            else:
                documents = self._load_text(path)

        except (OSError, ValueError, KeyError) as exc:
            raise DocumentLoadError(
                f"Failed to load '{path}': {exc}"
            ) from exc

        return documents

    def load_directory(
        self,
        directory_path: str | Path,
        recursive: bool = True,
    ) -> list[LoadedDocument]:
        """Load all supported documents from a directory."""
        directory = Path(directory_path)

        if not directory.exists():
            raise FileNotFoundError(
                f"Directory not found: {directory}"
            )

        if not directory.is_dir():
            raise ValueError(
                f"Expected a directory, got: {directory}"
            )

        pattern = "**/*" if recursive else "*"

        files = sorted(
            path
            for path in directory.glob(pattern)
            if path.is_file()
            and path.suffix.lower() in SUPPORTED_EXTENSIONS
        )

        documents = []

        for file_path in files:
            try:
                documents.extend(self.load_file(file_path))
            except DocumentLoadError as exc:
                raise DocumentLoadError(
                    f"Error processing directory file "
                    f"'{file_path}': {exc}"
                ) from exc

        return documents

    def _load_pdf(self, path: Path) -> list[LoadedDocument]:
        """Extract text from each non-empty PDF page."""
        reader = PdfReader(str(path))
        documents = []

        for page_number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()

            if not text:
                continue

            documents.append(
                LoadedDocument(
                    content=text,
                    metadata={
                        "source": str(path.resolve()),
                        "filename": path.name,
                        "file_type": "pdf",
                        "page": page_number,
                    },
                )
            )

        return documents

    def _load_docx(self, path: Path) -> list[LoadedDocument]:
        """Extract paragraphs and table text from a DOCX file."""
        doc = DocxDocument(str(path))

        sections = []

        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                sections.append(text)

        for table in doc.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip().replace("\n", " ")
                    for cell in row.cells
                ]
                row_text = " | ".join(
                    cell for cell in cells if cell
                )
                if row_text:
                    sections.append(row_text)

        content = "\n".join(sections).strip()

        if not content:
            return []

        return [
            LoadedDocument(
                content=content,
                metadata={
                    "source": str(path.resolve()),
                    "filename": path.name,
                    "file_type": "docx",
                },
            )
        ]

    def _load_text(self, path: Path) -> list[LoadedDocument]:
        """Read plain-text and Markdown files."""
        content = path.read_text(
            encoding="utf-8-sig",
            errors="replace",
        ).strip()

        if not content:
            return []

        return [
            LoadedDocument(
                content=content,
                metadata={
                    "source": str(path.resolve()),
                    "filename": path.name,
                    "file_type": path.suffix.lower().lstrip("."),
                },
            )
        ]